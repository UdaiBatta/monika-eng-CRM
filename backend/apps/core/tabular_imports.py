import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import serializers

MAX_IMPORT_ROWS = 500
MAX_UNCOMPRESSED_XLSX_BYTES = 25 * 1024 * 1024


class TabularImportUploadSerializer(serializers.Serializer):
    file = serializers.FileField()

    def validate_file(self, value):
        if value.size > 5 * 1024 * 1024:
            raise serializers.ValidationError("The import file must be 5 MB or smaller.")
        if not value.name.lower().endswith((".csv", ".xlsx")):
            raise serializers.ValidationError("Upload a .csv or .xlsx file.")
        return value


def _header(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _map_rows(header_values, rows, *, required_headers, allowed_headers, label):
    headers = [_header(value) for value in header_values]
    if not headers or not any(headers):
        raise serializers.ValidationError("The import file does not contain a header row.")
    if len(headers) != len(set(headers)):
        raise serializers.ValidationError("Column names must be unique.")
    missing = sorted(required_headers - set(headers))
    if missing:
        raise serializers.ValidationError(
            f"Missing required columns: {', '.join(missing)}. Download the template and try again."
        )
    unknown = sorted(set(headers) - allowed_headers)
    if unknown:
        raise serializers.ValidationError(
            f"Unknown columns: {', '.join(unknown)}. Download the template and try again."
        )

    result = []
    for values in rows:
        mapped = {name: _cell(value) for name, value in zip(headers, values, strict=False) if name}
        if not any(mapped.values()):
            continue
        result.append(mapped)
        if len(result) > MAX_IMPORT_ROWS:
            raise serializers.ValidationError(f"Import at most {MAX_IMPORT_ROWS} {label} at a time.")
    if not result:
        raise serializers.ValidationError(f"The import file contains no {label} rows.")
    return result


def parse_tabular_upload(upload, *, required_headers, allowed_headers, label):
    suffix = Path(upload.name).suffix.lower()
    data = upload.read()
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise serializers.ValidationError("CSV files must use UTF-8 text encoding.") from exc
        rows = csv.reader(io.StringIO(text))
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise serializers.ValidationError("The CSV file is empty.") from exc
        return _map_rows(
            headers,
            rows,
            required_headers=required_headers,
            allowed_headers=allowed_headers,
            label=label,
        )
    if suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > 250
                    or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_XLSX_BYTES
                ):
                    raise serializers.ValidationError(
                        "The Excel workbook expands beyond the safe import limit."
                    )
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration as exc:
                raise serializers.ValidationError("The Excel workbook is empty.") from exc
            return _map_rows(
                headers,
                rows,
                required_headers=required_headers,
                allowed_headers=allowed_headers,
                label=label,
            )
        except serializers.ValidationError:
            raise
        except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
            raise serializers.ValidationError("The uploaded Excel workbook is not valid.") from exc
    raise serializers.ValidationError("Upload a .csv or .xlsx file.")
